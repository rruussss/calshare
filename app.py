"""
CalShare - AI-Powered Calendar Sharing Platform
Accepts any file format and uses AI to extract calendar events
"""

import os
import json
import uuid
import sqlite3
import hashlib
import io
import base64
import re
from datetime import datetime, timedelta
from urllib.parse import quote as urlencode
from flask import (
    Flask, render_template, request, jsonify, redirect, 
    url_for, send_file, Response, abort
)
from werkzeug.utils import secure_filename
import qrcode
from icalendar import Calendar, Event
from dateutil import parser as date_parser
import pytz
import anthropic
from PIL import Image
import pytesseract
import fitz  # PyMuPDF for PDF handling

app = Flask(__name__)

# Add urlencode filter to Jinja2
@app.template_filter('urlencode')
def urlencode_filter(s):
    """URL encode filter for Jinja2 templates"""
    if isinstance(s, str):
        return urlencode(s)
    return s
app.config['SECRET_KEY'] = os.environ.get('SECRET_KEY', os.urandom(24).hex())
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max file size
app.config['DATABASE'] = os.environ.get('DATABASE_URL', 'calendars.db')
app.config['UPLOAD_FOLDER'] = 'uploads'

# Ensure upload folder exists
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

# Allowed file extensions - we accept almost anything
ALLOWED_EXTENSIONS = {
    'ics', 'ical',  # Calendar files
    'txt', 'csv', 'json',  # Text files
    'pdf',  # PDF files
    'png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp', 'tiff',  # Image files
    'doc', 'docx', 'rtf',  # Document files
    'xls', 'xlsx',  # Spreadsheet files
}

# Initialize Anthropic client (will fail gracefully if API key not set)
try:
    claude_client = anthropic.Anthropic(api_key=os.environ.get('ANTHROPIC_API_KEY'))
except Exception as e:
    print(f"Warning: Anthropic client initialization failed: {e}")
    claude_client = None

def get_db():
    """Get database connection"""
    conn = sqlite3.connect(app.config['DATABASE'])
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    """Initialize the database"""
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS calendars (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            slug TEXT UNIQUE NOT NULL,
            name TEXT NOT NULL,
            description TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            access_count INTEGER DEFAULT 0
        )
    ''')
    
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            calendar_id INTEGER NOT NULL,
            uid TEXT NOT NULL,
            title TEXT NOT NULL,
            description TEXT,
            location TEXT,
            start_time TIMESTAMP NOT NULL,
            end_time TIMESTAMP NOT NULL,
            all_day BOOLEAN DEFAULT FALSE,
            category TEXT DEFAULT 'general',
            FOREIGN KEY (calendar_id) REFERENCES calendars(id)
        )
    ''')
    
    conn.commit()
    conn.close()

def allowed_file(filename):
    """Check if file extension is allowed"""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def get_file_type(filename):
    """Determine the type of file"""
    ext = filename.rsplit('.', 1)[1].lower() if '.' in filename else ''
    
    if ext in {'ics', 'ical'}:
        return 'calendar'
    elif ext in {'png', 'jpg', 'jpeg', 'gif', 'bmp', 'webp', 'tiff'}:
        return 'image'
    elif ext == 'pdf':
        return 'pdf'
    elif ext in {'txt', 'csv', 'json', 'rtf'}:
        return 'text'
    elif ext in {'doc', 'docx'}:
        return 'word'
    elif ext in {'xls', 'xlsx'}:
        return 'excel'
    else:
        return 'unknown'

def generate_slug(name):
    """Generate a unique slug from the calendar name"""
    base_slug = ''.join(c if c.isalnum() else '-' for c in name.lower())
    base_slug = '-'.join(filter(None, base_slug.split('-')))[:30]
    unique_id = hashlib.md5(f"{name}{datetime.now().isoformat()}{uuid.uuid4()}".encode()).hexdigest()[:8]
    return f"{base_slug}-{unique_id}"

def extract_text_from_image(image_data):
    """Extract text from image using OCR"""
    try:
        if isinstance(image_data, bytes):
            image = Image.open(io.BytesIO(image_data))
        else:
            image = image_data
        
        # Convert to RGB if necessary
        if image.mode != 'RGB':
            image = image.convert('RGB')
        
        # Use pytesseract for OCR
        # Note: Tesseract must be installed on the system
        text = pytesseract.image_to_string(image)
        return text
    except Exception as e:
        print(f"OCR error: {e}")
        print("Note: Tesseract OCR may not be installed. OCR features will not work.")
        return ""

def extract_text_from_pdf(pdf_data):
    """Extract text from PDF, using OCR for scanned pages"""
    try:
        doc = fitz.open(stream=pdf_data, filetype="pdf")
        full_text = []
        
        for page_num in range(len(doc)):
            page = doc[page_num]
            
            # First try to extract text directly
            text = page.get_text()
            
            # If no text found, the page might be scanned - use OCR
            if not text.strip():
                # Render page to image
                mat = fitz.Matrix(2.0, 2.0)  # 2x zoom for better OCR
                pix = page.get_pixmap(matrix=mat)
                img_data = pix.tobytes("png")
                
                # OCR the image
                image = Image.open(io.BytesIO(img_data))
                text = pytesseract.image_to_string(image)
            
            full_text.append(text)
        
        doc.close()
        return "\n\n".join(full_text)
    except Exception as e:
        print(f"PDF extraction error: {e}")
        return ""

def parse_ics_file(file_content):
    """Parse ICS file and return list of events"""
    events = []
    try:
        cal = Calendar.from_ical(file_content)
        for component in cal.walk():
            if component.name == "VEVENT":
                event = {
                    'uid': str(component.get('uid', uuid.uuid4())),
                    'title': str(component.get('summary', 'Untitled Event')),
                    'description': str(component.get('description', '')),
                    'location': str(component.get('location', '')),
                    'category': 'general'
                }
                
                dtstart = component.get('dtstart')
                if dtstart:
                    start = dtstart.dt
                    if isinstance(start, datetime):
                        event['start_time'] = start.isoformat()
                        event['all_day'] = False
                    else:
                        event['start_time'] = datetime.combine(start, datetime.min.time()).isoformat()
                        event['all_day'] = True
                
                dtend = component.get('dtend')
                if dtend:
                    end = dtend.dt
                    if isinstance(end, datetime):
                        event['end_time'] = end.isoformat()
                    else:
                        event['end_time'] = datetime.combine(end, datetime.min.time()).isoformat()
                elif dtstart:
                    start = dtstart.dt
                    if isinstance(start, datetime):
                        event['end_time'] = (start + timedelta(hours=1)).isoformat()
                    else:
                        event['end_time'] = datetime.combine(start, datetime.min.time()).isoformat()
                
                events.append(event)
    except Exception as e:
        print(f"Error parsing ICS: {e}")
        raise ValueError(f"Failed to parse calendar file: {str(e)}")
    
    return events

def parse_with_ai(text_content, file_type="text", image_base64=None):
    """Use Claude AI to parse calendar events from any text or image"""
    
    # Check if Anthropic client is available
    if not claude_client:
        raise ValueError("Anthropic API key not configured. Please set ANTHROPIC_API_KEY environment variable.")
    
    current_year = datetime.now().year
    
    system_prompt = f"""You are a calendar event extraction expert. Your job is to extract calendar events from any format of input - text, OCR output, schedules, etc.

For each event you find, extract:
- title: The name/title of the event
- start_time: Start date and time in ISO format (YYYY-MM-DDTHH:MM:SS). If no year specified, assume {current_year} or {current_year + 1} (whichever makes more sense for future events)
- end_time: End date and time in ISO format. If not specified, assume 1 hour after start for timed events, or end of day for all-day events
- location: Where the event takes place (if mentioned)
- description: Any additional details about the event
- category: Categorize as one of: practice, game, meeting, event, deadline, general
- all_day: true if it's an all-day event, false otherwise

IMPORTANT RULES:
1. Be thorough - extract ALL events you can find
2. If times are ambiguous (like "9am"), make reasonable assumptions
3. If dates use formats like "12/1" or "Dec 1", parse them correctly
4. For recurring patterns (e.g., "every Tuesday"), list each individual occurrence if dates are given
5. Return ONLY valid JSON array, no other text
6. If you cannot find any events, return an empty array []

Return a JSON array of events like:
[
  {{
    "title": "Team Practice",
    "start_time": "2024-12-01T09:00:00",
    "end_time": "2024-12-01T11:00:00",
    "location": "Main Field",
    "description": "Weekly practice session",
    "category": "practice",
    "all_day": false
  }}
]"""

    try:
        messages = []
        
        # If we have an image, send it directly to Claude for vision analysis
        if image_base64:
            messages.append({
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": "image/png",
                            "data": image_base64
                        }
                    },
                    {
                        "type": "text",
                        "text": f"Extract all calendar events from this image. The image appears to be a {file_type}. Also consider this OCR text that was extracted: {text_content[:5000] if text_content else 'No OCR text available'}"
                    }
                ]
            })
        else:
            messages.append({
                "role": "user",
                "content": f"Extract all calendar events from the following {file_type} content:\n\n{text_content[:10000]}"
            })
        
        response = claude_client.messages.create(
            model="claude-sonnet-4-20250514",
            max_tokens=4096,
            system=system_prompt,
            messages=messages
        )
        
        response_text = response.content[0].text.strip()
        
        # Try to extract JSON from the response
        # Sometimes Claude wraps it in ```json blocks
        json_match = re.search(r'\[[\s\S]*\]', response_text)
        if json_match:
            events_json = json_match.group()
            events = json.loads(events_json)
        else:
            events = []
        
        # Add UIDs to events
        for event in events:
            event['uid'] = f"evt-{uuid.uuid4().hex[:12]}"
            
            # Validate and fix dates
            try:
                start = date_parser.parse(event['start_time'])
                event['start_time'] = start.isoformat()
            except:
                pass
            
            try:
                end = date_parser.parse(event['end_time'])
                event['end_time'] = end.isoformat()
            except:
                # Default to 1 hour after start
                try:
                    start = date_parser.parse(event['start_time'])
                    event['end_time'] = (start + timedelta(hours=1)).isoformat()
                except:
                    pass
        
        return events
        
    except Exception as e:
        print(f"AI parsing error: {e}")
        raise ValueError(f"Failed to parse events with AI: {str(e)}")

def process_uploaded_file(file):
    """Process any uploaded file and extract calendar events"""
    filename = secure_filename(file.filename)
    file_type = get_file_type(filename)
    file_content = file.read()
    
    events = []
    
    if file_type == 'calendar':
        # Standard ICS file - parse directly
        events = parse_ics_file(file_content)
        
    elif file_type == 'image':
        # Image file - use OCR + AI
        ocr_text = extract_text_from_image(file_content)
        
        # Also send the image directly to Claude for vision analysis
        image_base64 = base64.b64encode(file_content).decode('utf-8')
        events = parse_with_ai(ocr_text, "image/schedule", image_base64)
        
    elif file_type == 'pdf':
        # PDF file - extract text (with OCR if needed) + AI
        pdf_text = extract_text_from_pdf(file_content)
        
        # Also try to get first page as image for Claude vision
        try:
            doc = fitz.open(stream=file_content, filetype="pdf")
            page = doc[0]
            mat = fitz.Matrix(2.0, 2.0)
            pix = page.get_pixmap(matrix=mat)
            img_data = pix.tobytes("png")
            image_base64 = base64.b64encode(img_data).decode('utf-8')
            doc.close()
        except:
            image_base64 = None
        
        events = parse_with_ai(pdf_text, "PDF document", image_base64)
        
    elif file_type == 'text':
        # Text file - direct AI parsing
        text_content = file_content.decode('utf-8', errors='ignore')
        events = parse_with_ai(text_content, "text file")
        
    elif file_type == 'word':
        # Word document - extract text and use AI
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_content))
            text_content = "\n".join([para.text for para in doc.paragraphs])
        except:
            text_content = file_content.decode('utf-8', errors='ignore')
        events = parse_with_ai(text_content, "Word document")
        
    elif file_type == 'excel':
        # Excel file - extract and use AI
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_content))
            text_parts = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join([str(cell) if cell else "" for cell in row])
                    if row_text.strip():
                        text_parts.append(row_text)
            text_content = "\n".join(text_parts)
        except:
            text_content = ""
        events = parse_with_ai(text_content, "Excel spreadsheet")
        
    else:
        # Unknown format - try as text
        try:
            text_content = file_content.decode('utf-8', errors='ignore')
            events = parse_with_ai(text_content, "document")
        except:
            raise ValueError("Unable to process this file format")
    
    return events

def create_ics_file(events, calendar_name):
    """Create an ICS file from events"""
    cal = Calendar()
    cal.add('prodid', '-//CalShare//calshare.app//')
    cal.add('version', '2.0')
    cal.add('calscale', 'GREGORIAN')
    cal.add('method', 'PUBLISH')
    cal.add('x-wr-calname', calendar_name)
    
    for event_data in events:
        event = Event()
        event.add('uid', event_data['uid'])
        event.add('summary', event_data['title'])
        
        if event_data.get('description'):
            event.add('description', event_data['description'])
        if event_data.get('location'):
            event.add('location', event_data['location'])
        
        start = date_parser.parse(event_data['start_time'])
        end = date_parser.parse(event_data['end_time'])
        
        if event_data.get('all_day'):
            event.add('dtstart', start.date())
            event.add('dtend', end.date())
        else:
            event.add('dtstart', start)
            event.add('dtend', end)
        
        event.add('dtstamp', datetime.now(pytz.UTC))
        cal.add_component(event)
    
    return cal.to_ical()

def generate_qr_code(url):
    """Generate QR code as base64 image"""
    qr = qrcode.QRCode(
        version=1,
        error_correction=qrcode.constants.ERROR_CORRECT_L,
        box_size=10,
        border=4,
    )
    qr.add_data(url)
    qr.make(fit=True)
    
    img = qr.make_image(fill_color="black", back_color="white")
    buffer = io.BytesIO()
    img.save(buffer, format='PNG')
    buffer.seek(0)
    
    return base64.b64encode(buffer.getvalue()).decode()

# ============ ROUTES ============

@app.route('/')
def index():
    """Home page"""
    return render_template('index.html')

@app.route('/create', methods=['GET', 'POST'])
def create_calendar():
    """Create a new calendar"""
    if request.method == 'GET':
        return render_template('create.html')
    
    try:
        data = request.get_json()
        
        if not data:
            return jsonify({'error': 'No data provided'}), 400
        
        name = data.get('name', '').strip()
        description = data.get('description', '').strip()
        events = data.get('events', [])
        custom_slug = data.get('custom_slug', '').strip()
        
        if not name:
            return jsonify({'error': 'Calendar name is required'}), 400
        
        if not events:
            return jsonify({'error': 'At least one event is required'}), 400
        
        if custom_slug:
            slug = ''.join(c if c.isalnum() or c == '-' else '' for c in custom_slug.lower())
        else:
            slug = generate_slug(name)
        
        conn = get_db()
        cursor = conn.cursor()
        cursor.execute('SELECT id FROM calendars WHERE slug = ?', (slug,))
        if cursor.fetchone():
            if custom_slug:
                conn.close()
                return jsonify({'error': 'This custom URL is already taken'}), 400
            slug = generate_slug(name)
        
        cursor.execute(
            'INSERT INTO calendars (slug, name, description) VALUES (?, ?, ?)',
            (slug, name, description)
        )
        calendar_id = cursor.lastrowid
        
        for event in events:
            cursor.execute('''
                INSERT INTO events (calendar_id, uid, title, description, location, 
                                   start_time, end_time, all_day, category)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                calendar_id,
                event.get('uid', str(uuid.uuid4())),
                event.get('title', 'Untitled Event'),
                event.get('description', ''),
                event.get('location', ''),
                event.get('start_time'),
                event.get('end_time'),
                event.get('all_day', False),
                event.get('category', 'general')
            ))
        
        conn.commit()
        conn.close()
        
        return jsonify({
            'success': True,
            'slug': slug,
            'url': url_for('view_calendar', slug=slug, _external=True)
        })
        
    except Exception as e:
        print(f"Error creating calendar: {e}")
        return jsonify({'error': str(e)}), 500

@app.route('/upload', methods=['POST'])
def upload_file():
    """Upload and parse any file type"""
    if 'file' not in request.files:
        return jsonify({'error': 'No file provided'}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({'error': 'No file selected'}), 400
    
    try:
        events = process_uploaded_file(file)
        
        if not events:
            return jsonify({
                'success': True,
                'events': [],
                'message': 'No calendar events could be extracted from this file. Please try a different file or enter events manually.'
            })
        
        return jsonify({
            'success': True,
            'events': events,
            'message': f'Successfully extracted {len(events)} events'
        })
        
    except Exception as e:
        print(f"Upload error: {e}")
        return jsonify({'error': str(e)}), 400

@app.route('/parse-text', methods=['POST'])
def parse_text():
    """Parse calendar events from pasted text"""
    try:
        data = request.get_json()
        text = data.get('text', '')
        
        if not text.strip():
            return jsonify({'error': 'No text provided'}), 400
        
        events = parse_with_ai(text, "pasted text")
        
        return jsonify({
            'success': True,
            'events': events,
            'message': f'Successfully extracted {len(events)} events'
        })
        
    except Exception as e:
        print(f"Parse error: {e}")
        return jsonify({'error': str(e)}), 400

@app.route('/c/<slug>')
def view_calendar(slug):
    """View a shared calendar"""
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('SELECT * FROM calendars WHERE slug = ?', (slug,))
    calendar = cursor.fetchone()
    
    if not calendar:
        conn.close()
        abort(404)
    
    cursor.execute('UPDATE calendars SET access_count = access_count + 1 WHERE slug = ?', (slug,))
    conn.commit()
    
    cursor.execute('SELECT * FROM events WHERE calendar_id = ? ORDER BY start_time', (calendar['id'],))
    events = [dict(row) for row in cursor.fetchall()]
    
    conn.close()
    
    categories = list(set(e['category'] for e in events if e['category']))
    
    base_url = request.url_root.rstrip('/')
    ics_url = f"{base_url}/c/{slug}/download.ics"
    
    user_agent = request.headers.get('User-Agent', '').lower()
    device_info = {
        'is_ios': 'iphone' in user_agent or 'ipad' in user_agent,
        'is_android': 'android' in user_agent,
        'is_mac': 'macintosh' in user_agent,
        'is_windows': 'windows' in user_agent,
        'is_mobile': any(x in user_agent for x in ['iphone', 'ipad', 'android', 'mobile'])
    }
    
    calendar_url = f"{base_url}/c/{slug}"
    qr_code = generate_qr_code(calendar_url)
    
    return render_template('view.html',
        calendar=dict(calendar),
        events=events,
        categories=categories,
        ics_url=ics_url,
        calendar_url=calendar_url,
        qr_code=qr_code,
        device_info=device_info
    )

@app.route('/c/<slug>/download.ics')
def download_ics(slug):
    """Download ICS file for calendar"""
    categories = request.args.getlist('category')
    
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('SELECT * FROM calendars WHERE slug = ?', (slug,))
    calendar = cursor.fetchone()
    
    if not calendar:
        conn.close()
        abort(404)
    
    if categories:
        placeholders = ','.join('?' * len(categories))
        cursor.execute(f'''
            SELECT * FROM events WHERE calendar_id = ? AND category IN ({placeholders})
            ORDER BY start_time
        ''', [calendar['id']] + categories)
    else:
        cursor.execute('SELECT * FROM events WHERE calendar_id = ? ORDER BY start_time', (calendar['id'],))
    
    events = [dict(row) for row in cursor.fetchall()]
    conn.close()
    
    ics_content = create_ics_file(events, calendar['name'])
    
    response = Response(
        ics_content,
        mimetype='text/calendar',
        headers={'Content-Disposition': f'attachment; filename="{calendar["name"]}.ics"'}
    )
    
    return response

@app.route('/c/<slug>/qr')
def get_qr_code(slug):
    """Get QR code image for calendar"""
    conn = get_db()
    cursor = conn.cursor()
    cursor.execute('SELECT id FROM calendars WHERE slug = ?', (slug,))
    calendar = cursor.fetchone()
    conn.close()
    
    if not calendar:
        abort(404)
    
    base_url = request.url_root.rstrip('/')
    calendar_url = f"{base_url}/c/{slug}"
    
    qr = qrcode.QRCode(version=1, error_correction=qrcode.constants.ERROR_CORRECT_L, box_size=10, border=4)
    qr.add_data(calendar_url)
    qr.make(fit=True)
    
    img = qr.make_image(fill_color="black", back_color="white")
    buffer = io.BytesIO()
    img.save(buffer, format='PNG')
    buffer.seek(0)
    
    return send_file(buffer, mimetype='image/png')

@app.route('/c/<slug>/google')
def google_calendar_link(slug):
    """Generate Google Calendar add link for a single event"""
    event_id = request.args.get('event_id')
    
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('SELECT * FROM calendars WHERE slug = ?', (slug,))
    calendar = cursor.fetchone()
    
    if not calendar:
        conn.close()
        abort(404)
    
    if event_id:
        cursor.execute('SELECT * FROM events WHERE calendar_id = ? AND id = ?', (calendar['id'], event_id))
    else:
        cursor.execute('SELECT * FROM events WHERE calendar_id = ? ORDER BY start_time LIMIT 1', (calendar['id'],))
    
    event = cursor.fetchone()
    conn.close()
    
    if not event:
        abort(404)
    
    start = date_parser.parse(event['start_time'])
    end = date_parser.parse(event['end_time'])
    
    if event['all_day']:
        date_format = '%Y%m%d'
        dates = f"{start.strftime(date_format)}/{end.strftime(date_format)}"
    else:
        date_format = '%Y%m%dT%H%M%S'
        dates = f"{start.strftime(date_format)}/{end.strftime(date_format)}"
    
    params = f"action=TEMPLATE&text={urlencode(event['title'])}&dates={dates}"
    if event['description']:
        params += f"&details={urlencode(event['description'])}"
    if event['location']:
        params += f"&location={urlencode(event['location'])}"
    
    return redirect(f"https://calendar.google.com/calendar/render?{params}")

@app.route('/api/calendars/<slug>')
def api_get_calendar(slug):
    """API endpoint to get calendar data"""
    conn = get_db()
    cursor = conn.cursor()
    
    cursor.execute('SELECT * FROM calendars WHERE slug = ?', (slug,))
    calendar = cursor.fetchone()
    
    if not calendar:
        conn.close()
        return jsonify({'error': 'Calendar not found'}), 404
    
    cursor.execute('SELECT * FROM events WHERE calendar_id = ? ORDER BY start_time', (calendar['id'],))
    events = [dict(row) for row in cursor.fetchall()]
    
    conn.close()
    
    return jsonify({'calendar': dict(calendar), 'events': events})

@app.errorhandler(404)
def not_found(e):
    return render_template('404.html'), 404

@app.errorhandler(500)
def server_error(e):
    return render_template('500.html'), 500

if __name__ == '__main__':
    init_db()
    port = int(os.environ.get('PORT', 5000))
    app.run(debug=False, host='0.0.0.0', port=port)
